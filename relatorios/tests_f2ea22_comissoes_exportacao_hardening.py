from io import BytesIO
from decimal import Decimal
from openpyxl import load_workbook
from django.urls import reverse
from django.contrib.auth import get_user_model
from relatorios.tests_f2ea22_comissoes_escopo_db import F2EA22ComissoesEscopoDBTests
from configuracoes.models import ComissaoVendedor

class F2EA22ExportacoesHardeningTests(F2EA22ComissoesEscopoDBTests):
    def xlsx(self, params=None):
        self.client.force_login(self.admin)
        r=self.client.get(reverse("relatorios:exportar_comissoes_xlsx"),params or {})
        self.assertEqual(r.status_code,200)
        return load_workbook(BytesIO(r.content),data_only=True)["Comissoes"]

    def test_filtro_competencia_exporta_somente_snapshot_filtrado(self):
        ws=self.xlsx({"ano":"2026","mes":"10"})
        nomes=[r[5] for r in ws.iter_rows(min_row=2,values_only=True)]
        self.assertEqual(nomes,["Vendedor Loja A2"])

    def test_exportacao_usa_nome_historico_do_snapshot(self):
        ws=self.xlsx({"loja":self.a1.pk})
        nomes=[r[5] for r in ws.iter_rows(min_row=2,values_only=True)]
        self.assertIn("Vendedor Snapshot Historico",nomes)
        self.assertNotIn("Nome Atual",nomes)

    def test_xlsx_campos_monetarios_e_percentual_sao_numericos(self):
        ws=self.xlsx({"loja":self.a1.pk})
        row=list(ws.iter_rows(min_row=2,values_only=True))[0]
        for idx in (2,3,4,6,8):
            self.assertIsInstance(row[idx],(int,float))

    def test_exportacao_nao_limita_a_50_registros(self):
        U=get_user_model()
        for i in range(55):
            vendedor=U.objects.create_user(
                username=f"extra_f2ea22_{i:02d}",password="x",
                first_name="Extra",last_name=f"{i:02d}",
                matriz=self.m1,ativo=True,
            )
            ComissaoVendedor.objects.create(
                fechamento=self.f1,vendedor=vendedor,
                vendedor_uuid_original=vendedor.uuid,
                vendedor_nome_original=f"Extra Snapshot {i:02d}",
                total_vendas_vendedor=Decimal("10.00"),
                elegivel=True,valor_comissao=Decimal("1.00"),
            )
        ws=self.xlsx({"loja":self.a1.pk})
        self.assertEqual(ws.max_row-1,56)

    def test_get_exportacoes_nao_altera_snapshots(self):
        before=ComissaoVendedor.objects.count()
        self.client.force_login(self.admin)
        self.client.get(reverse("relatorios:exportar_comissoes_pdf"))
        self.client.get(reverse("relatorios:exportar_comissoes_xlsx"))
        self.assertEqual(ComissaoVendedor.objects.count(),before)