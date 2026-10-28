from decimal import Decimal
from django.contrib.auth.decorators import login_required
from django.views.decorators.http import require_GET
from django.shortcuts import render
from django.http import HttpResponse
from django.core.paginator import Paginator
from django.db.models import Sum
from django.utils import timezone
from django.core.exceptions import PermissionDenied
from empresas.models import Loja, Matriz
from clientes.models import Cliente
from vouchers.models import Voucher
from vouchers.selectors import get_relatorio_usos_voucher
from cashback.selectors import get_relatorio_movimentacoes_cashback

from core.services import get_contexto_operacional_usuario

from .selectors import get_dashboard_resumo, get_relatorio_clientes, get_relatorio_aniversariantes

from accounts.decorators import require_permission
from accounts.permissions import PERMISSAO_RELATORIOS_DASHBOARD
from compras.selectors import get_fornecedores, get_pedidos_compra
from compras.models import Fornecedor, PedidoCompra
from configuracoes.models import FechamentoComissao, ComissaoVendedor
from financeiro.models import CentroCusto
from .services_dre import apurar_dre, apurar_evolucao_mensal, REGIME_COMPETENCIA


@login_required
@require_permission(PERMISSAO_RELATORIOS_DASHBOARD)
def hub_cadastros(request):
    return render(request, 'relatorios/cadastros.html')

@login_required
@require_permission(PERMISSAO_RELATORIOS_DASHBOARD)
def hub_vendas(request):
    return render(request, 'relatorios/vendas.html')


@login_required
@require_permission(PERMISSAO_RELATORIOS_DASHBOARD)
def hub_financeiro(request):
    pode_acessar_dre = request.user.is_superuser or getattr(request.user, "perfil", None) != "operador"
    return render(
        request,
        "relatorios/financeiro.html",
        {"pode_acessar_dre": pode_acessar_dre},
    )


def hub_estoque(request):
    return render(request, 'relatorios/estoque.html')

@login_required
@require_permission(PERMISSAO_RELATORIOS_DASHBOARD)
def dashboard(request):

    contexto = get_contexto_operacional_usuario(
        request.user
    )

    resumo = get_dashboard_resumo(
        matriz=contexto['matriz']
    )

    return render(
        request,
        'relatorios/dashboard.html',
        {
            'resumo': resumo,
        }
    )

@login_required
@require_permission(PERMISSAO_RELATORIOS_DASHBOARD)
def relatorio_clientes(request):
    contexto = get_contexto_operacional_usuario(request.user)
    busca = (request.GET.get('q') or '').strip()
    ativo_raw = (request.GET.get('ativo') or '').strip()
    tipo_pessoa = (request.GET.get('tipo_pessoa') or '').strip()
    loja_raw = (request.GET.get('loja') or '').strip()
    ativo = True if ativo_raw == '1' else False if ativo_raw == '0' else None
    loja_id = int(loja_raw) if loja_raw.isdigit() else None
    clientes_qs = get_relatorio_clientes(
        matriz=contexto['matriz'], ativo=ativo,
        tipo_pessoa=tipo_pessoa or None, loja_id=loja_id, busca=busca or None,
    )
    clientes = Paginator(clientes_qs, 50).get_page(request.GET.get('page'))
    lojas = Loja.objects.filter(matriz=contexto['matriz']).order_by('nome', 'id')
    query_params = request.GET.copy()
    query_params.pop('page', None)
    query_string = query_params.urlencode()
    return render(request, 'relatorios/clientes.html', {
        'clientes': clientes,
        'busca': busca,
        'ativo_filtro': ativo_raw,
        'tipo_pessoa_filtro': tipo_pessoa,
        'loja_filtro': loja_raw,
        'lojas': lojas,
        'query_string': query_string,
    })

@login_required
@require_permission(PERMISSAO_RELATORIOS_DASHBOARD)
def relatorio_aniversariantes(request):
    contexto = get_contexto_operacional_usuario(request.user)
    mes_atual = timezone.localdate().month
    try:
        mes = int(request.GET.get('mes') or mes_atual)
    except (TypeError, ValueError):
        mes = mes_atual
    if mes < 1 or mes > 12:
        mes = mes_atual
    qs = get_relatorio_aniversariantes(matriz=contexto['matriz'], mes=mes)
    paginator = Paginator(qs, 50)
    page_obj = paginator.get_page(request.GET.get('page'))
    return render(request, 'relatorios/aniversariantes.html', {
        'page_obj': page_obj,
        'mes_filtro': mes,
        'meses': [(i, nome) for i, nome in enumerate(('Janeiro','Fevereiro','Março','Abril','Maio','Junho','Julho','Agosto','Setembro','Outubro','Novembro','Dezembro'), 1)],
    })


@login_required
@require_permission(PERMISSAO_RELATORIOS_DASHBOARD)
def hub_vouchers(request):
    return render(request, 'relatorios/vouchers.html')


@login_required
@require_permission(PERMISSAO_RELATORIOS_DASHBOARD)
def relatorio_utilizacao_vouchers(request):
    contexto = get_contexto_operacional_usuario(request.user)
    matriz = contexto['matriz']
    perfil = getattr(request.user, 'perfil', None)
    operador = perfil == 'operador'

    if operador:
        lojas_autorizadas = request.user.lojas.filter(matriz=matriz).order_by('nome', 'id')
        loja_id = request.session.get('loja_operacional_id')
        loja = lojas_autorizadas.filter(pk=loja_id).first() if loja_id else None
        if loja is None:
            loja = lojas_autorizadas.first()
            if loja is not None:
                request.session['loja_operacional_id'] = loja.pk
        lojas = lojas_autorizadas.filter(pk=loja.pk) if loja else lojas_autorizadas.none()
    else:
        lojas = Loja.objects.filter(matriz=matriz).order_by('nome', 'id')
        loja_raw = (request.GET.get('loja') or '').strip()
        loja = lojas.filter(pk=loja_raw).first() if loja_raw.isdigit() else None

    data_inicio_raw = (request.GET.get('data_inicio') or '').strip()
    data_fim_raw = (request.GET.get('data_fim') or '').strip()
    voucher_raw = (request.GET.get('voucher') or '').strip()

    from datetime import date
    def parse_data(valor):
        try:
            return date.fromisoformat(valor) if valor else None
        except ValueError:
            return None

    usos = get_relatorio_usos_voucher(
        matriz=matriz,
        loja=loja,
        data_inicio=parse_data(data_inicio_raw),
        data_fim=parse_data(data_fim_raw),
        voucher_id=voucher_raw if voucher_raw.isdigit() else None,
    )
    totais = usos.aggregate(
        valor_compras=Sum('valor_compra'),
        valor_descontos=Sum('valor_desconto'),
    )
    quantidade = usos.count()
    page_obj = Paginator(usos, 50).get_page(request.GET.get('page'))
    vouchers = Voucher.objects.filter(matriz=matriz).order_by('nome', 'id')

    return render(request, 'relatorios/vouchers_utilizacao.html', {
        'page_obj': page_obj,
        'quantidade': quantidade,
        'valor_compras': totais['valor_compras'] or 0,
        'valor_descontos': totais['valor_descontos'] or 0,
        'lojas': lojas,
        'loja_filtro': loja,
        'operador': operador,
        'vouchers': vouchers,
        'voucher_filtro': voucher_raw,
        'data_inicio': data_inicio_raw,
        'data_fim': data_fim_raw,
    })


@login_required
@require_permission(PERMISSAO_RELATORIOS_DASHBOARD)
def hub_cashback(request):
    return render(request, 'relatorios/cashback.html')


@login_required
@require_permission(PERMISSAO_RELATORIOS_DASHBOARD)
def relatorio_movimentacoes_cashback(request):
    contexto = get_contexto_operacional_usuario(request.user)
    matriz = contexto['matriz']
    operador = getattr(request.user, 'perfil', None) == 'operador'

    if operador:
        lojas_autorizadas = request.user.lojas.filter(matriz=matriz).order_by('nome', 'id')
        loja_id = request.session.get('loja_operacional_id')
        loja = lojas_autorizadas.filter(pk=loja_id).first() if loja_id else None
        if loja is None:
            loja = lojas_autorizadas.first()
            if loja is not None:
                request.session['loja_operacional_id'] = loja.pk
        lojas = lojas_autorizadas.filter(pk=loja.pk) if loja else lojas_autorizadas.none()
    else:
        lojas = Loja.objects.filter(matriz=matriz).order_by('nome', 'id')
        loja_raw = (request.GET.get('loja') or '').strip()
        loja = lojas.filter(pk=loja_raw).first() if loja_raw.isdigit() else None

    data_inicio = (request.GET.get('data_inicio') or '').strip()
    data_fim = (request.GET.get('data_fim') or '').strip()
    tipo = (request.GET.get('tipo') or '').strip().upper()
    if tipo not in ('ENTRADA', 'SAIDA'):
        tipo = ''
    cliente_raw = (request.GET.get('cliente') or '').strip()
    cliente_id = int(cliente_raw) if cliente_raw.isdigit() else None

    resultado = get_relatorio_movimentacoes_cashback(
        matriz=matriz,
        loja=loja,
        cliente_id=cliente_id,
        data_inicio=data_inicio or None,
        data_fim=data_fim or None,
        tipo=tipo or None,
    )
    page_obj = Paginator(resultado['movimentacoes'], 50).get_page(request.GET.get('page'))
    clientes = Cliente.objects.filter(matriz=matriz).order_by('nome', 'id')

    return render(request, 'relatorios/cashback_movimentacoes.html', {
        'page_obj': page_obj,
        'quantidade': resultado['quantidade_movimentacoes'],
        'total_gerado': resultado['total_gerado'],
        'total_utilizado': resultado['total_utilizado'],
        'lojas': lojas,
        'loja_filtro': loja,
        'operador': operador,
        'clientes': clientes,
        'cliente_filtro': cliente_raw,
        'tipo_filtro': tipo,
        'data_inicio': data_inicio,
        'data_fim': data_fim,
    })


def _bloquear_operador_compras(request):
    if getattr(request.user, 'perfil', None) == 'operador':
        raise PermissionDenied('Relatorios de compras e fornecedores sao restritos ao perfil administrativo da matriz.')
@login_required
@require_permission(PERMISSAO_RELATORIOS_DASHBOARD)
def hub_compras(request):
    _bloquear_operador_compras(request)
    return render(request, 'relatorios/compras.html')

@login_required
@require_permission(PERMISSAO_RELATORIOS_DASHBOARD)
def relatorio_pedidos_compra(request):
    _bloquear_operador_compras(request)
    c=get_contexto_operacional_usuario(request.user);q=(request.GET.get('q') or '').strip();st=(request.GET.get('status') or '').strip();data_inicio=(request.GET.get('data_inicio') or '').strip();data_fim=(request.GET.get('data_fim') or '').strip();qs=get_pedidos_compra(matriz=c['matriz'],busca=q,status=st)
    from datetime import date
    try:
        di=date.fromisoformat(data_inicio) if data_inicio else None
    except ValueError:
        di=None
    try:
        df=date.fromisoformat(data_fim) if data_fim else None
    except ValueError:
        df=None
    qs=qs.filter(data_emissao__gte=di) if di else qs;qs=qs.filter(data_emissao__lte=df) if df else qs;xs=list(qs);qp=request.GET.copy();qp.pop('page',None)
    return render(request,'relatorios/compras_pedidos.html',{'page_obj':Paginator(xs,50).get_page(request.GET.get('page')),'quantidade':len(xs),'total_pedidos':sum((x.total for x in xs),Decimal('0.00')),'busca':q,'status_filtro':st,'data_inicio':data_inicio,'data_fim':data_fim,'status_choices':PedidoCompra._meta.get_field('status').choices,'query_string':qp.urlencode()})

@login_required
@require_permission(PERMISSAO_RELATORIOS_DASHBOARD)
def relatorio_fornecedores(request):
    _bloquear_operador_compras(request)
    c=get_contexto_operacional_usuario(request.user);q=(request.GET.get('q') or '').strip();st=(request.GET.get('status') or '').strip();qs=get_fornecedores(matriz=c['matriz'],busca=q,status=st);n=qs.count();qp=request.GET.copy();qp.pop('page',None)
    return render(request,'relatorios/compras_fornecedores.html',{'page_obj':Paginator(qs,50).get_page(request.GET.get('page')),'quantidade':n,'busca':q,'status_filtro':st,'status_choices':Fornecedor._meta.get_field('status').choices,'query_string':qp.urlencode()})


@login_required
@require_permission(PERMISSAO_RELATORIOS_DASHBOARD)
def hub_caixa(request):
    return render(request, "relatorios/caixa.html")

@login_required
@require_permission(PERMISSAO_RELATORIOS_DASHBOARD)
def hub_comissoes(request):
    return render(request, 'relatorios/comissoes.html')

def _dados_relatorio_comissoes(request):
    contexto=get_contexto_operacional_usuario(request.user);matriz=contexto['matriz'];operador=getattr(request.user,'perfil',None)==getattr(request.user,'PERFIL_OPERADOR','operador')
    loja_raw=(request.GET.get("loja") or '').strip();ano_raw=(request.GET.get('ano') or '').strip();mes_raw=(request.GET.get('mes') or '').strip()
    if operador:
        lojas_autorizadas=request.user.lojas.filter(matriz=matriz).order_by('nome','id');loja_operacional_id=request.session.get('loja_operacional_id');loja=lojas_autorizadas.filter(pk=loja_operacional_id).first() if loja_operacional_id else None;lojas=lojas_autorizadas.filter(pk=loja.pk) if loja else lojas_autorizadas.none()
    else:
        lojas=Loja.objects.filter(matriz=matriz).order_by('nome','id');loja=lojas.filter(pk=loja_raw).first() if loja_raw.isdigit() else None
    fechamentos=FechamentoComissao.objects.filter(matriz=matriz,loja__in=lojas).select_related('loja').order_by('-competencia_ano','-competencia_mes','loja__nome','id')
    if loja is not None:fechamentos=fechamentos.filter(loja=loja)
    if ano_raw.isdigit():fechamentos=fechamentos.filter(competencia_ano=int(ano_raw))
    if mes_raw.isdigit() and 1<=int(mes_raw)<=12:fechamentos=fechamentos.filter(competencia_mes=int(mes_raw))
    comissoes=ComissaoVendedor.objects.filter(fechamento__in=fechamentos).select_related('fechamento','fechamento__loja','vendedor').order_by('-fechamento__competencia_ano','-fechamento__competencia_mes','fechamento__loja__nome','vendedor_nome_original','id')
    return matriz,operador,lojas,loja,ano_raw,mes_raw,fechamentos,comissoes

@login_required
@require_GET
@require_permission(PERMISSAO_RELATORIOS_DASHBOARD)
def relatorio_comissoes(request):
    matriz,operador,lojas,loja,ano_raw,mes_raw,fechamentos,comissoes=_dados_relatorio_comissoes(request)
    page_obj=Paginator(comissoes,50).get_page(request.GET.get('page'));qp=request.GET.copy();qp.pop('page',None)
    return render(request,'relatorios/relatorio_comissoes.html',{'page_obj':page_obj,'fechamentos':fechamentos,'lojas':lojas,'loja_filtro':loja,'operador':operador,'ano_filtro':ano_raw,'mes_filtro':mes_raw,'query_string':qp.urlencode()})

def _formatar_brl(valor):
    return f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _resumo_exportacao_comissoes(dados):
    matriz,operador,lojas,loja,ano_raw,mes_raw,fechamentos,comissoes=dados
    fechamentos=list(fechamentos)
    comissoes=list(comissoes)
    total_vendido=sum((f.total_vendas_loja for f in fechamentos), Decimal("0"))
    total_comissoes=sum((c.valor_comissao for c in comissoes), Decimal("0"))
    elegiveis=sum(1 for c in comissoes if c.elegivel)
    return {"fechamentos":fechamentos,"comissoes":comissoes,"total_vendido":total_vendido,"total_comissoes":total_comissoes,"elegiveis":elegiveis,"nao_elegiveis":len(comissoes)-elegiveis}

@login_required
@require_GET
@require_permission(PERMISSAO_RELATORIOS_DASHBOARD)
def exportar_comissoes_pdf(request):
    from io import BytesIO
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER,TA_RIGHT
    from reportlab.lib.pagesizes import A4,landscape
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate,Table,TableStyle,Paragraph,Spacer
    dados=_dados_relatorio_comissoes(request);resumo=_resumo_exportacao_comissoes(dados)
    matriz,operador,lojas,loja,ano_raw,mes_raw,fechamentos,comissoes=dados
    buffer=BytesIO();doc=SimpleDocTemplate(buffer,pagesize=landscape(A4),rightMargin=10*mm,leftMargin=10*mm,topMargin=10*mm,bottomMargin=10*mm)
    styles=getSampleStyleSheet()
    title=ParagraphStyle("Titulo",parent=styles["Title"],fontName="Helvetica-Bold",fontSize=16,leading=19,alignment=TA_CENTER,spaceAfter=4*mm)
    meta=ParagraphStyle("Meta",parent=styles["Normal"],fontSize=8.5,leading=11)
    right=ParagraphStyle("Right",parent=meta,alignment=TA_RIGHT)
    story=[Paragraph("Relatório de Comissões",title)]
    ano=ano_raw or "Todos";mes=mes_raw
    competencia=f"{int(mes):02d}/{ano}" if mes.isdigit() and ano!="Todos" else (f"Todos/{ano}" if ano!="Todos" else "Todas")
    loja_txt=str(loja) if loja is not None else ("Loja operacional" if operador else "Todas as lojas")
    info=Table([[Paragraph(f"<b>Matriz:</b> {matriz}",meta),Paragraph(f"<b>Loja:</b> {loja_txt}",meta)],[Paragraph(f"<b>Competência:</b> {competencia}",meta),Paragraph(f"<b>Emitido em:</b> {timezone.localtime().strftime('%d/%m/%Y %H:%M')}",right)]],colWidths=[135*mm,135*mm])
    story += [info,Spacer(1,4*mm)]
    rt=Table([["Total vendido","Comissões","Elegíveis","Não elegíveis","Registros"],[_formatar_brl(resumo["total_vendido"]),_formatar_brl(resumo["total_comissoes"]),str(resumo["elegiveis"]),str(resumo["nao_elegiveis"]),str(len(resumo["comissoes"]))]],colWidths=[54*mm]*5)
    rt.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#E9ECEF")),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("ALIGN",(0,0),(-1,-1),"CENTER"),("GRID",(0,0),(-1,-1),0.4,colors.HexColor("#CED4DA")),("FONTSIZE",(0,0),(-1,-1),8.5)]))
    story += [rt,Spacer(1,4*mm)]
    rows=[["Competência","Loja","Total vendido","Meta","Percentual","Vendedor","Vendas vendedor","Elegível","Comissão"]]
    for c in resumo["comissoes"]:
        f=c.fechamento
        rows.append([f"{f.competencia_mes:02d}/{f.competencia_ano}",f.loja.nome,_formatar_brl(f.total_vendas_loja),"" if f.valor_meta_atingida is None else _formatar_brl(f.valor_meta_atingida),f"{f.percentual_aplicado:.2f}%",c.vendedor_nome_original,_formatar_brl(c.total_vendas_vendedor),"Sim" if c.elegivel else "Não",_formatar_brl(c.valor_comissao)])
    table=Table(rows,repeatRows=1,colWidths=[20*mm,34*mm,27*mm,24*mm,22*mm,43*mm,30*mm,18*mm,25*mm])
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#343A40")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,-1),7.5),("GRID",(0,0),(-1,-1),0.3,colors.HexColor("#DEE2E6")),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("ALIGN",(2,1),(4,-1),"RIGHT"),("ALIGN",(6,1),(6,-1),"RIGHT"),("ALIGN",(8,1),(8,-1),"RIGHT"),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F8F9FA")])]))
    story.append(table);doc.build(story)
    response=HttpResponse(buffer.getvalue(),content_type="application/pdf");response["Content-Disposition"]='attachment; filename="relatorio_comissoes.pdf"';return response

@login_required
@require_GET
@require_permission(PERMISSAO_RELATORIOS_DASHBOARD)
def exportar_comissoes_xlsx(request):
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    from openpyxl.utils import get_column_letter
    dados=_dados_relatorio_comissoes(request);resumo=_resumo_exportacao_comissoes(dados)
    matriz,operador,lojas,loja,ano_raw,mes_raw,fechamentos,comissoes=dados
    wb=Workbook();ws=wb.active;ws.title="Comissoes"
    ws.append(["Competencia","Loja","Total vendido","Meta","Percentual","Vendedor","Vendas do vendedor","Elegivel","Comissao"])
    for c in resumo["comissoes"]:
        f=c.fechamento
        ws.append([f"{f.competencia_mes:02d}/{f.competencia_ano}",f.loja.nome,float(f.total_vendas_loja),None if f.valor_meta_atingida is None else float(f.valor_meta_atingida),float(f.percentual_aplicado)/100,c.vendedor_nome_original,float(c.total_vendas_vendedor),"Sim" if c.elegivel else "Não",float(c.valor_comissao)])
    fill=PatternFill("solid",fgColor="343A40")
    for cell in ws[1]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=fill;cell.alignment=Alignment(horizontal="center")
    thin=Side(style="thin",color="D9D9D9")
    for row in ws.iter_rows():
        for cell in row:cell.border=Border(bottom=thin)
    for row in range(2,ws.max_row+1):
        for col in (3,4,7,9):ws.cell(row,col).number_format='R$ #,##0.00'
        ws.cell(row,5).number_format='0.00%'
    for i,w in enumerate([14,28,18,18,14,32,20,12,18],1):ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes="A2";ws.auto_filter.ref=f"A1:I{ws.max_row}"
    rs=wb.create_sheet("Resumo");rs.append(["Relatório de Comissões",""]);rs.append(["Matriz",str(matriz)]);rs.append(["Loja",str(loja) if loja is not None else ("Loja operacional" if operador else "Todas as lojas")])
    ano=ano_raw or "Todos";mes=mes_raw;comp=f"{int(mes):02d}/{ano}" if mes.isdigit() and ano!="Todos" else (f"Todos/{ano}" if ano!="Todos" else "Todas")
    rs.append(["Competência",comp]);rs.append(["Emitido em",timezone.localtime().strftime("%d/%m/%Y %H:%M")]);rs.append([]);rs.append(["Indicador","Valor"]);rs.append(["Total vendido",float(resumo["total_vendido"])]);rs.append(["Total de comissões",float(resumo["total_comissoes"])]);rs.append(["Vendedores elegíveis",resumo["elegiveis"]]);rs.append(["Vendedores não elegíveis",resumo["nao_elegiveis"]]);rs.append(["Registros",len(resumo["comissoes"])])
    rs["A1"].font=Font(bold=True,size=16);rs.merge_cells("A1:B1")
    for cell in rs[7]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=fill
    rs["B8"].number_format='R$ #,##0.00';rs["B9"].number_format='R$ #,##0.00';rs.column_dimensions["A"].width=28;rs.column_dimensions["B"].width=42
    out=BytesIO();wb.save(out)
    response=HttpResponse(out.getvalue(),content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet");response["Content-Disposition"]='attachment; filename="relatorio_comissoes.xlsx"';return response

def _periodo_anterior_dre(inicio, fim):
    from datetime import date, timedelta
    def ultimo_dia_mes(valor):
        proximo = date(valor.year + (1 if valor.month == 12 else 0), 1 if valor.month == 12 else valor.month + 1, 1)
        return proximo - timedelta(days=1)
    mes_calendario_completo = inicio.day == 1 and fim == ultimo_dia_mes(inicio) and inicio.year == fim.year and inicio.month == fim.month
    if mes_calendario_completo:
        anterior_fim = inicio - timedelta(days=1)
        anterior_inicio = anterior_fim.replace(day=1)
        return anterior_inicio, anterior_fim
    dias = (fim - inicio).days
    anterior_fim = inicio - timedelta(days=1)
    anterior_inicio = anterior_fim - timedelta(days=dias)
    return anterior_inicio, anterior_fim


def _contexto_dre(request):
    from datetime import date, timedelta

    if request.user.is_superuser:
        matrizes = Matriz.objects.order_by("nome", "id")
        matriz_raw = (request.GET.get("matriz") or "").strip()
        matriz = matrizes.filter(pk=matriz_raw).first() if matriz_raw.isdigit() else None
        operador = False
    else:
        contexto = get_contexto_operacional_usuario(request.user)
        matriz = contexto["matriz"]
        matrizes = Matriz.objects.none()
        matriz_raw = ""
        operador = getattr(request.user, "perfil", None) == getattr(request.user, "PERFIL_OPERADOR", "operador")

    hoje = timezone.localdate()
    inicio_padrao = hoje.replace(day=1)
    fim_padrao = (inicio_padrao.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
    data_inicio_input = (request.GET.get("data_inicio") or "").strip()
    data_fim_input = (request.GET.get("data_fim") or "").strip()
    erro_filtros = []

    def parse_data(raw, padrao):
        if not raw:
            return padrao
        try:
            return date.fromisoformat(raw)
        except ValueError:
            erro_filtros.append("Data inválida")
            return None

    data_inicio = parse_data(data_inicio_input, inicio_padrao)
    data_fim = parse_data(data_fim_input, fim_padrao)
    if data_inicio is not None and data_fim is not None and data_fim < data_inicio:
        erro_filtros.append("Período inválido")

    if request.user.is_superuser and matriz is None:
        erro_filtros.append("Selecione uma matriz")

    filtros_validos = not erro_filtros
    loja_raw = (request.GET.get("loja") or "").strip()
    centro_raw = (request.GET.get("centro_custo") or "").strip()

    if matriz is None:
        lojas = Loja.objects.none()
        loja = None
        centros_custo = CentroCusto.objects.none()
        centro_custo = None
    elif operador:
        autorizadas = request.user.lojas.filter(matriz=matriz).order_by("nome", "id")
        lid = request.session.get("loja_operacional_id")
        loja = autorizadas.filter(pk=lid).first() if lid else autorizadas.first()
        if loja and not lid:
            request.session["loja_operacional_id"] = loja.pk
        lojas = autorizadas.filter(pk=loja.pk) if loja else autorizadas.none()
        centros_custo = CentroCusto.objects.filter(matriz=matriz,ativo=True).order_by("nome", "id")
        centro_custo = centros_custo.filter(pk=centro_raw).first() if centro_raw.isdigit() else None
    else:
        lojas = Loja.objects.filter(matriz=matriz).order_by("nome", "id")
        loja = lojas.filter(pk=loja_raw).first() if loja_raw.isdigit() else None
        centros_custo = CentroCusto.objects.filter(matriz=matriz,ativo=True).order_by("nome", "id")
        centro_custo = centros_custo.filter(pk=centro_raw).first() if centro_raw.isdigit() else None

    periodo_anterior = None
    apuracao = None
    evolucao = []
    if not erro_filtros:
        periodo_anterior = _periodo_anterior_dre(data_inicio, data_fim)
        apuracao = apurar_dre(
            matriz=matriz, inicio=data_inicio, fim=data_fim, loja=loja,
            centro_custo=centro_custo, periodo_anterior=periodo_anterior,
        )
        evolucao = apurar_evolucao_mensal(
            matriz=matriz, inicio=data_inicio, fim=data_fim, loja=loja,
            centro_custo=centro_custo,
        )

    return {
        "matriz": matriz, "matrizes": matrizes, "matriz_filtro": matriz_raw,
        "operador": operador, "lojas": lojas, "loja_filtro": loja,
        "centros_custo": centros_custo, "centro_custo_filtro": centro_custo,
        "data_inicio": data_inicio, "data_fim": data_fim,
        "data_inicio_input": data_inicio_input or inicio_padrao.isoformat(),
        "data_fim_input": data_fim_input or fim_padrao.isoformat(),
        "periodo_anterior": periodo_anterior, "erro_filtros": erro_filtros,
        "filtros_validos": filtros_validos, "regime": REGIME_COMPETENCIA,
        "resultado": apuracao["resultado"] if apuracao else None,
        "analise_vertical": apuracao["analise_vertical"] if apuracao else {},
        "drilldown": apuracao["drilldown"] if apuracao else {},
        "comparativo": apuracao.get("comparativo") if apuracao else None,
        "evolucao": evolucao, "query_string": request.GET.urlencode(),
    }

def _bloquear_operador_dre(request):
    if not request.user.is_superuser and getattr(request.user, "perfil", None) == "operador":
        raise PermissionDenied("DRE gerencial restrita ao perfil administrativo da matriz.")


@login_required
@require_GET
@require_permission(PERMISSAO_RELATORIOS_DASHBOARD)
def dre(request):
    _bloquear_operador_dre(request)
    return render(request,"relatorios/dre.html",_contexto_dre(request))

@login_required
@require_GET
@require_permission(PERMISSAO_RELATORIOS_DASHBOARD)
def exportar_dre_pdf(request):
    _bloquear_operador_dre(request)
    from io import BytesIO
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    contexto = _contexto_dre(request)
    if not contexto.get("filtros_validos", contexto.get("resultado") is not None):
        return HttpResponse(
            "; ".join(contexto.get("erro_filtros") or ["Filtros inválidos"]),
            status=400,
            content_type="text/plain; charset=utf-8",
        )
    resultado = contexto["resultado"]; av = contexto["analise_vertical"]
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=14*mm, leftMargin=14*mm, topMargin=12*mm, bottomMargin=12*mm)
    styles=getSampleStyleSheet(); titulo=ParagraphStyle("DRETitulo",parent=styles["Title"],fontName="Helvetica-Bold",fontSize=15,leading=18)
    story=[Paragraph("Demonstração do Resultado - DRE",titulo),Paragraph(f"Matriz: {contexto['matriz']} | Período: {contexto['data_inicio'].strftime('%d/%m/%Y')} a {contexto['data_fim'].strftime('%d/%m/%Y')} | Regime: {contexto['regime']}",styles["Normal"]),Spacer(1,5*mm)]
    linhas=[("Receita Bruta","receita_bruta"),("(-) Deduções da Receita","deducoes_receita"),("(=) Receita Líquida","receita_liquida"),("(-) CMV","cmv"),("(=) Lucro Bruto","lucro_bruto"),("(-) Despesas Operacionais","despesas_operacionais"),("(=) Resultado Operacional","resultado_operacional"),("(+/-) Resultado Financeiro","resultado_financeiro"),("(+/-) Outras Receitas / Despesas","outras_receitas_despesas"),("(=) Resultado Antes dos Impostos","resultado_antes_impostos")]
    rows=[["Linha","Valor","Análise vertical"]]+[[rotulo,_formatar_brl(getattr(resultado,campo)),f"{av.get(campo,0):.2f}%"] for rotulo,campo in linhas]
    table=Table(rows,repeatRows=1,colWidths=[92*mm,48*mm,38*mm]);table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#343A40")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("GRID",(0,0),(-1,-1),0.3,colors.HexColor("#DEE2E6")),("ALIGN",(1,1),(-1,-1),"RIGHT"),("FONTSIZE",(0,0),(-1,-1),9)]))
    story.append(table);doc.build(story)
    response=HttpResponse(buffer.getvalue(),content_type="application/pdf");response["Content-Disposition"]='attachment; filename="relatorio_dre.pdf"';return response

@login_required
@require_GET
@require_permission(PERMISSAO_RELATORIOS_DASHBOARD)
def exportar_dre_xlsx(request):
    _bloquear_operador_dre(request)
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    contexto = _contexto_dre(request)
    if not contexto.get("filtros_validos", contexto.get("resultado") is not None):
        return HttpResponse(
            "; ".join(contexto.get("erro_filtros") or ["Filtros inválidos"]),
            status=400,
            content_type="text/plain; charset=utf-8",
        )
    resultado = contexto["resultado"]; av = contexto["analise_vertical"]
    wb=Workbook();ws=wb.active;ws.title="DRE";ws.append(["Demonstração do Resultado - DRE","",""]);ws.merge_cells("A1:C1");ws["A1"].font=Font(bold=True,size=16)
    ws.append(["Matriz",str(contexto["matriz"]),""]);ws.append(["Período",f"{contexto['data_inicio'].strftime('%d/%m/%Y')} a {contexto['data_fim'].strftime('%d/%m/%Y')}",""]);ws.append(["Regime",contexto["regime"],""]);ws.append([]);ws.append(["Linha","Valor","Análise vertical"])
    linhas=[("Receita Bruta","receita_bruta"),("(-) Deduções da Receita","deducoes_receita"),("(=) Receita Líquida","receita_liquida"),("(-) CMV","cmv"),("(=) Lucro Bruto","lucro_bruto"),("(-) Despesas Operacionais","despesas_operacionais"),("(=) Resultado Operacional","resultado_operacional"),("(+/-) Resultado Financeiro","resultado_financeiro"),("(+/-) Outras Receitas / Despesas","outras_receitas_despesas"),("(=) Resultado Antes dos Impostos","resultado_antes_impostos")]
    for rotulo,campo in linhas: ws.append([rotulo,float(getattr(resultado,campo)),float(av.get(campo,0))/100])
    fill=PatternFill("solid",fgColor="343A40")
    for c in ws[6]: c.font=Font(bold=True,color="FFFFFF");c.fill=fill;c.alignment=Alignment(horizontal="center")
    for row in range(7,ws.max_row+1): ws.cell(row,2).number_format='R$ #,##0.00';ws.cell(row,3).number_format='0.00%'
    ws.column_dimensions["A"].width=38;ws.column_dimensions["B"].width=22;ws.column_dimensions["C"].width=20;ws.freeze_panes="A7"
    out=BytesIO();wb.save(out)
    response=HttpResponse(out.getvalue(),content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet");response["Content-Disposition"]='attachment; filename="relatorio_dre.xlsx"';return response
