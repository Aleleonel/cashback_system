from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from relatorios.services_dre import ResultadoDRE

ZERO=0
class DREHTTPExportacoesTests(TestCase):
    def setUp(self):
        U=get_user_model()
        self.user=U.objects.create_superuser(username="f25d37_admin",email="f25d37@example.test",password="TesteF25D37!")
        self.client.force_login(self.user)
        self.resultado=ResultadoDRE(
            receita_bruta=210,acrescimos=0,deducoes_receita=20,receita_liquida=190,cmv=80,lucro_bruto=110,
            despesas_operacionais=40,resultado_operacional=70,resultado_financeiro=5,outras_receitas_despesas=0,
            resultado_antes_impostos=75,margem_bruta=0,margem_operacional=0,memoria_calculo={}
        )
        self.ctx={"matriz":"Matriz Teste","operador":False,"lojas":[],"loja_filtro":None,"centros_custo":[],"centro_custo_filtro":None,
                  "data_inicio":__import__("datetime").date(2026,9,1),"data_fim":__import__("datetime").date(2026,9,30),
                  "regime":"COMPETENCIA","resultado":self.resultado,
                  "analise_vertical":{k:0 for k in ("receita_bruta","deducoes_receita","receita_liquida","cmv","lucro_bruto","despesas_operacionais","resultado_operacional","resultado_financeiro","outras_receitas_despesas","resultado_antes_impostos")},
                  "drilldown":{},"comparativo":None,"evolucao":[],"query_string":"data_inicio=2026-09-01&data_fim=2026-09-30"}

    @patch("relatorios.views._contexto_dre")
    def test_pdf_http_real(self,m):
        m.return_value=self.ctx
        r=self.client.get(reverse("relatorios:exportar_dre_pdf"))
        self.assertEqual(r.status_code,200);self.assertEqual(r["Content-Type"],"application/pdf")
        self.assertIn("relatorio_dre.pdf",r["Content-Disposition"]);self.assertTrue(r.content.startswith(b"%PDF"))
        self.assertGreater(len(r.content),500)

    @patch("relatorios.views._contexto_dre")
    def test_xlsx_http_real_e_abre_workbook(self,m):
        m.return_value=self.ctx
        r=self.client.get(reverse("relatorios:exportar_dre_xlsx"))
        self.assertEqual(r.status_code,200)
        self.assertEqual(r["Content-Type"],"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("relatorio_dre.xlsx",r["Content-Disposition"])
        wb=load_workbook(BytesIO(r.content),data_only=False)
        self.assertIn("DRE",wb.sheetnames);ws=wb["DRE"]
        self.assertEqual(ws["A1"].value,"Demonstração do Resultado - DRE")
        self.assertEqual(ws["B7"].value,210.0)
        self.assertEqual(ws["B16"].value,75.0)

    def test_endpoints_exigem_autenticacao(self):
        self.client.logout()
        for nome in ("relatorios:dre","relatorios:exportar_dre_pdf","relatorios:exportar_dre_xlsx"):
            r=self.client.get(reverse(nome));self.assertIn(r.status_code,(302,403))