from io import BytesIO
from openpyxl import load_workbook
from django.test import TestCase
from django.urls import reverse

from relatorios.tests_f2ea22_comissoes_escopo_db import F2EA22ComissoesEscopoDBTests


class F2EA22ExportacoesFuncionaisTests(F2EA22ComissoesEscopoDBTests):
    """Reutiliza as fixtures multiempresa ja validadas no contrato F2EA22."""

    def _login_admin(self):
        self.client.force_login(self.admin)

    def test_pdf_responde_arquivo_valido(self):
        self._login_admin()
        r=self.client.get(reverse("relatorios:exportar_comissoes_pdf"))
        self.assertEqual(r.status_code,200)
        self.assertEqual(r["Content-Type"],"application/pdf")
        self.assertTrue(r.content.startswith(b"%PDF"))
        self.assertIn("attachment;",r["Content-Disposition"])

    def test_xlsx_responde_arquivo_valido_e_colunas(self):
        self._login_admin()
        r=self.client.get(reverse("relatorios:exportar_comissoes_xlsx"))
        self.assertEqual(r.status_code,200)
        wb=load_workbook(BytesIO(r.content),data_only=True)
        ws=wb["Comissoes"]
        self.assertEqual([c.value for c in ws[1]],[
            "Competencia","Loja","Total vendido","Meta","Percentual",
            "Vendedor","Vendas do vendedor","Elegivel","Comissao",
        ])
        self.assertGreaterEqual(ws.max_row,2)

    def test_xlsx_nao_vaza_matriz_externa(self):
        self._login_admin()
        r=self.client.get(reverse("relatorios:exportar_comissoes_xlsx"))
        ws=load_workbook(BytesIO(r.content),data_only=True)["Comissoes"]
        valores=[" ".join("" if v is None else str(v) for v in row) for row in ws.iter_rows(values_only=True)]
        texto="\n".join(valores)
        self.assertNotIn(self.x.nome,texto)

    def test_xlsx_filtro_loja_propria(self):
        self._login_admin()
        r=self.client.get(reverse("relatorios:exportar_comissoes_xlsx"),{"loja":self.a1.pk})
        ws=load_workbook(BytesIO(r.content),data_only=True)["Comissoes"]
        lojas={row[1] for row in ws.iter_rows(min_row=2,values_only=True)}
        self.assertTrue(lojas <= {self.a1.nome})

    def test_operador_nao_consegue_forcar_loja_por_querystring_xlsx(self):
        self.client.force_login(self.op)
        session=self.client.session;session["loja_operacional_id"]=self.a1.pk;session.save()
        r=self.client.get(reverse("relatorios:exportar_comissoes_xlsx"),{"loja":self.a2.pk})
        ws=load_workbook(BytesIO(r.content),data_only=True)["Comissoes"]
        lojas={row[1] for row in ws.iter_rows(min_row=2,values_only=True)}
        self.assertNotIn(self.a2.nome,lojas)
        self.assertTrue(lojas <= {self.a1.nome})

    def test_operador_sem_loja_operacional_exporta_apenas_cabecalho(self):
        self.client.force_login(self.op)
        session=self.client.session;session.pop("loja_operacional_id",None);session.save()
        r=self.client.get(reverse("relatorios:exportar_comissoes_xlsx"))
        ws=load_workbook(BytesIO(r.content),data_only=True)["Comissoes"]
        self.assertEqual(ws.max_row,1)

    def test_post_exportacoes_rejeitado(self):
        self._login_admin()
        for nome in ("relatorios:exportar_comissoes_pdf","relatorios:exportar_comissoes_xlsx"):
            self.assertEqual(self.client.post(reverse(nome)).status_code,405)